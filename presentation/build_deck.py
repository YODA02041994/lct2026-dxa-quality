#!/usr/bin/env python3
"""Сборка презентации к сдаче (ЛЦТ-2026, задача № 4).

Слайды 7–11 шаблона организатора заполняются как есть (сетка и оформление не меняются); остальные слайды
строятся в рамке шаблона (фон, логотипы, розовая плашка раздела), а внутри белой карточки — по правилам
научной презентации: чёрный текст Arial, заголовок-тезис 24 пт, цветные графики, ссылка на источник 10 пт.

    python3 presentation/make_figures.py && python3 presentation/build_deck.py
Шаблон: docs/materials/ЛЦТ2026_Шаблон_презентации.pptx (в git не хранится, scripts/download_data.sh).
Личные данные участников берутся из presentation/team.json (в git не хранится); без файла ставятся пометки «[…]».
"""
import copy
import json
import os
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.dirname(HERE)
sys.path.insert(0, os.path.expanduser("~/.claude/skills/science-deck/scripts"))
from pptx_kit import Deck, BLACK, GREY  # noqa: E402

TEMPLATE = os.path.join(ROOT, "docs", "materials", "ЛЦТ2026_Шаблон_презентации.pptx")
FIG, IMG = os.path.join(HERE, "fig"), os.path.join(HERE, "img")
OUT = os.path.join(HERE, "DXA-QC_презентация.pptx")
PINK, WHITE = RGBColor(0xFF, 0x00, 0x53), RGBColor(255, 255, 255)
TEAM = json.load(open(os.path.join(HERE, "team.json"), encoding="utf-8")) if os.path.exists(os.path.join(HERE, "team.json")) else {}

# ───────────────────────── 1. шаблон: оставить только обязательные слайды 7–11 ─────────────────────────
prs = Presentation(TEMPLATE)
ids = prs.slides._sldIdLst
for i, sld in reversed(list(enumerate(list(ids), 1))):
    if i not in (7, 8, 9, 10, 11):
        prs.part.drop_rel(sld.rId); ids.remove(sld)
tmp = os.path.join(HERE, "_trimmed.pptx"); prs.save(tmp)


class CardDeck(Deck):
    """Свободные слайды в рамке шаблона: плашка раздела + белая карточка с содержимым по правилам science-deck."""

    def __init__(self, template):
        super().__init__(template=template)
        self.SLIDE_H = self.prs.slide_height
        self.card = (Inches(0.38), Inches(1.22), Inches(12.57), Inches(5.72))
        self.L = Inches(0.78); self.R = Inches(0.78); self.CW = self.W - self.L - self.R
        self.TOP = Inches(1.45); self.H = self.card[1] + self.card[3] + Inches(0.02)      # «низ слайда» = низ карточки
        self.blank = next(l for l in self.prs.slide_layouts if l.name == "Пустой слайд")
        self.n = len(self.prs.slides)

    def slide(self, title=None, section="", number=True):
        s = self.prs.slides.add_slide(self.blank); self.n += 1
        for ph in list(s.placeholders):                                   # дата/колонтитул шаблона не нужны
            ph._element.getparent().remove(ph._element)
        pw = Inches(0.55 + 0.155 * len(section))
        pill = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.38), Inches(0.35), pw, Inches(0.68)); pill.adjustments[0] = 0.18
        pill.fill.solid(); pill.fill.fore_color.rgb = PINK; pill.line.fill.background(); pill.shadow.inherit = False
        tf = pill.text_frame; tf.margin_left = Inches(0.2); tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        r = tf.paragraphs[0].add_run(); r.text = section.upper(); r.font.name, r.font.size, r.font.bold, r.font.color.rgb = "Arial", Pt(16), True, WHITE
        c = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, *self.card); c.adjustments[0] = 0.035
        c.fill.solid(); c.fill.fore_color.rgb = WHITE; c.line.fill.background(); c.shadow.inherit = False
        if number:
            self.box(s, self.W - Inches(1.0), self.SLIDE_H - Inches(0.50), Inches(0.6), Inches(0.3), str(self.n), 11, False, WHITE, PP_ALIGN.RIGHT)
        if not title:
            return s, self.TOP
        h = Emu(int(self.lines(title, 24, True, self.CW) * 24 * 1.22 * 12700))
        self.box(s, self.L, self.TOP, self.CW, h, title, 24, True)
        self.rule(s, self.TOP + h + Inches(0.08))
        return s, self.TOP + h + Inches(0.24)


d = CardDeck(tmp)
S = list(d.prs.slides)


def shape(slide, sid):
    return next(sh for sh in slide.shapes if sh.shape_id == sid)


def set_par(par, text):
    """Заменить текст абзаца, сохранив оформление первого фрагмента шаблона."""
    runs = list(par.runs)
    if not runs:
        par.add_run().text = text; return
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def set_lines(sh, lines, size=None):
    tf = sh.text_frame; pars = list(tf.paragraphs)
    for p in pars[1:]:
        p._p.getparent().remove(p._p)
    set_par(pars[0], lines[0])
    for t in lines[1:]:
        np_ = copy.deepcopy(pars[0]._p); pars[0]._p.getparent().append(np_)
        from pptx.text.text import _Paragraph
        set_par(_Paragraph(np_, pars[0]._parent), t)
    if size:
        for p in tf.paragraphs:
            for r in p.runs:
                r.font.size = Pt(size)


def T(key, default):
    return TEAM.get(key, default)


def put_picture(slide, sid, path):
    """Вставить картинку в местозаполнитель шаблона, сохранив его положение и размер (иначе берётся геометрия макета)."""
    ph = shape(slide, sid); l, t, w, h = ph.left, ph.top, ph.width, ph.height
    pic = ph.insert_picture(path)
    pic.left, pic.top, pic.width, pic.height = l, t, w, h
    return pic


# ───────────────────────── 2. обязательные слайды 7–11 ─────────────────────────
s7 = S[0]
set_lines(shape(s7, 3), ["КОНТРОЛЬ КАЧЕСТВА ДЕНСИТОМЕТРИИ"])
set_lines(shape(s7, 5), [f"Команда «{T('team_name', '[название команды]')}»", "Задача № 4 · Департамент здравоохранения города Москвы"])
put_picture(s7, 4, os.path.join(IMG, "logo.png"))

s8 = S[1]
set_lines(shape(s8, 18), ["DXA-QC: автоматический контроль укладки"])
for par in shape(s8, 18).text_frame.paragraphs:                      # заголовок стоит на белой панели — тёмный цвет шаблона
    for r in par.runs:
        r.font.color.rgb = RGBColor(0x31, 0x0F, 0x53)
set_lines(shape(s8, 5), ["Сервис принимает DICOM-исследования денситометрии, определяет область (поясничный отдел позвоночника, "
                         "проксимальный отдел бедра), находит пять типов нарушений укладки и возвращает таблицу в формате заказчика, "
                         "снимок с разметкой нарушения и DICOM SR с текстом заключения."], 13)
set_lines(shape(s8, 8), ["Каждое решение подкреплено измерением: угол оси, запас поля в сантиметрах, рамка постороннего предмета. "
                         "Лаборант получает совет по стандарту ISCD, врач — отметку «проверить» для снимков у порога. "
                         "Детектор предметов обучен на 9000 открытых рентгенограмм грудной клетки и перенесён на денситометрию."], 13)
team_box = shape(s8, 14); pars = list(team_box.text_frame.paragraphs)
pars[0].runs[1].text = T("captain", "Семенов Андрей, детский ортопед, ассистент кафедры детской хирургии")
pars[1].runs[1].text = f"{T('n_members', '[число]')} человек"
set_par(pars[3], T("how_formed", "команда собрана 13.09.2026 в Пироговском Университете"))
set_par(pars[4], T("workplaces", "РНИМУ им. Н. И. Пирогова, РМАНПО"))
pars[5].runs[0].text = "Город и регион: "
val = copy.deepcopy(pars[0].runs[1]._r); val.text = T("city", "Москва"); pars[5]._p.append(val)
put_picture(s8, 2, os.path.join(IMG, "ov_hip_wide.png"))

s9 = S[2]
set_lines(shape(s9, 7), ["СОСТАВ КОМАНДЫ"])
members = T("members", [
    {"name": "Андрей Семенов", "role": "капитан; модель и клиническая экспертиза", "nick": "@drsav1994", "phone": "[телефон]", "work": "РНИМУ им. Н. И. Пирогова, кафедра детской хирургии", "avatar": "av_as.png"},
    {"name": "Александр [фамилия]", "role": "разметка ориентиров на снимках", "nick": "[ник]", "phone": "[телефон]", "work": "[место работы/учёбы]", "avatar": "av_a.png"},
    {"name": "Саак Серобян", "role": "[роль]", "nick": "[ник]", "phone": "[телефон]", "work": "РМАНПО, ординатура по кардиологии", "avatar": "av_ss.png"},
])
cards = [  # (рамка, подпись-данные, подпись-имя, фото) по shape_id шаблона, слева направо
    (17, 9, 15, 2), (56, 57, 58, 3), (59, 60, 61, 4), (62, 63, 64, 5), (65, 66, 67, 6)]
for i, (frame, info, name, photo) in enumerate(cards):
    if i < len(members):
        m = members[i]
        set_lines(shape(s9, name), [m["name"]])
        set_lines(shape(s9, info), [m["role"], m["nick"], m["phone"], m["work"]])
        av = os.path.join(IMG, m.get("avatar", ""))
        if os.path.isfile(av):
            put_picture(s9, photo, av)
    else:                                                              # лишние карточки удаляются — это разрешено правилами
        for sid in (frame, info, name, photo):
            sh = shape(s9, sid); sh._element.getparent().remove(sh._element)

s10 = S[3]
set_lines(shape(s10, 7), ["О КОМАНДЕ"])
set_lines(shape(s10, 37), [T("history", "Команда собрана 13.09.2026 в Пироговском Университете из врачей и разработчика. "
                                        "Вместе участвуем впервые; разметку ориентиров на 252 снимках сделали сами за четыре дня.")], 12)
set_lines(shape(s10, 43), [T("why", "В команде практикующие врачи: ошибка укладки при денситометрии меняет минеральную плотность и диагноз. "
                                    "Задача проверяется измерением, поэтому решение можно объяснить коллеге-рентгенологу.")], 12)
set_lines(shape(s10, 40), [T("challenges", "На критерий приходится 6–36 снимков с нарушением, поэтому сеть «по всей картинке» не обучалась. "
                                           "Мы разметили анатомические ориентиры, перешли к измерениям по снимку и перенесли детектор предметов "
                                           "с открытого набора рентгенограмм грудной клетки. ROC-AUC вырос с 0,74 до 0,85.")], 12)

s11 = S[4]
set_lines(shape(s11, 3), [
    "Вход: папка или zip с DICOM; дубликаты серий считаются один раз.",
    "Сеть ставит анатомические ориентиры (ошибка 2–5 мм); по снимку измеряются ось, запасы поля в сантиметрах, полоса таза, тонкие яркие линии.",
    "Ансамбль свёрточных сетей оценивает ротацию бедра и посторонние предметы; логистическая регрессия даёт вероятность по каждому из 5 критериев.",
    "Выход: results.xlsx в формате ТЗ, PNG-серия с разметкой, DICOM SR, HTTP-сервис и веб-страница; контейнер работает без GPU и без сети.",
    "ROC-AUC 0,849; F1 0,762; macro-F1 0,653 на снимках вне обучения."], 13)
set_lines(shape(s11, 7), [
    "Аудит качества денситометрии по кабинетам и лаборантам в ЕРИС: сортировка по вероятности нарушения.",
    "Подсказка лаборанту до ухода пациента: что исправить в укладке и нужно ли переснять.",
    "Допуск к ИИ-сервисам расчёта плотности только снимков, прошедших контроль.",
    "Развитие: сверка спорных меток с экспертами, контур бедра по плотной разметке, дообучение на снимках других аппаратов."], 13)

# ───────────────────────── 3. содержательные слайды ─────────────────────────
def figure_slide(section, title, fig, note, refs, max_h=3.35):
    s, y = d.slide(title, section)
    ph = d.pic(s, fig, y, max_w=d.CW, max_h=Inches(max_h))
    if note:
        d.box(s, d.L, d.below(ph, 0.10), d.CW, Inches(0.7), note, 18)
    d.refs(s, refs)
    return s


figure_slide("Задача", "Технические ошибки денситометрии найдены у 90 % пациентов, у 13 % они меняют интерпретацию",
             os.path.join(FIG, "lit_errors.png"),
             "Поворот бедра на 10° значимо сдвигает минеральную плотность шейки у 8–12 % обследованных.",
             "Krueger D, et al. J Clin Densitom. 2019;22(1):115–124. PMID: 30327243. DOI: 10.1016/j.jocd.2018.07.014 · "
             "Jung EY, et al. Arch Osteoporos. 2020;15(1):115. PMID: 32705454. DOI: 10.1007/s11657-020-00791-8 · "
             "Lekamwasam S, Lenora RS. J Clin Densitom. 2003;6(4):331–336. PMID: 14716045. DOI: 10.1385/jcd:6:4:331", 3.05)

s, y = d.slide("Пять типов нарушений в двух областях; на одном снимке их может быть несколько", "Таксономия")
d.table(s, [["Область", "Тип нарушения (формулировка заказчика)", "Что требует стандарт", "Что измеряет сервис"],
            ["Позвоночник", "Некорректная укладка", "в кадре Th12 и гребни таза", "полоса таза у нижнего края"],
            ["Позвоночник", "Не выровнена ось позвоночника", "столб вертикален, по центру", "наклон оси по массе кости, °"],
            ["Позвоночник", "Присутствуют посторонние предметы", "нет металла и фурнитуры", "тонкие яркие линии, карта предметов"],
            ["Бедро", "Некорректная укладка", "малый вертел виден минимально", "ось диафиза, °; сеть по вырезке вертела"],
            ["Бедро", "Некорректная область интереса", "головка, вертел, диафиз в поле", "запасы до краёв кадра, см"]],
        y + Inches(0.05), [1.7, 4.3, 3.0, 2.75], size=14, row_h_in=0.5)
d.box(s, d.L, y + Inches(3.25), d.CW, Inches(0.7), "Вероятности критериев независимы; итог по снимку: 1 − ∏(1 − pᵢ). Сколиоз, перелом и эндопротез нарушением не считаются.", 18)
d.refs(s, "ISCD Official Positions 2023, https://iscd.org/official-positions-2023/ · Albano D, et al. Acad Radiol. 2021;28(9):1272–1286. PMID: 32839098. DOI: 10.1016/j.acra.2020.07.028 · "
          "Разъяснения постановщика задачи (ЦДиТ ДЗМ), сессия 16.09.2026")

figure_slide("Данные", "На критерий приходится от 6 до 36 снимков с нарушением из 252 уникальных",
             os.path.join(FIG, "data.png"),
             "100 исследований, 499 файлов, 252 уникальных снимка. Проверка — 5 блоков по исследованиям: снимки одного пациента "
             "не попадают в обучение и проверку одновременно.",
             "Набор данных постановщика задачи (GE Lunar Prodigy Advance, DICOM CR 8 бит); истина — колонки критериев экспертной разметки", 2.85)

s, y = d.slide("При сотнях снимков работает схема «измерить — классифицировать — объяснить»", "Подход")
bw, bh, gap = Inches(2.72), Inches(1.75), Inches(0.31); x0 = d.L; yy = y + Inches(0.15)
boxes = [[[("1. Найти", True)], ["ориентиры: сеть"], ["ResNet18 + тепловые карты"]],
         [[("2. Измерить", True)], ["угол оси, запасы, см"], ["полоса таза, линии"]],
         [[("3. Классифицировать", True)], ["логистическая регрессия"], ["+ ансамбль сетей"]],
         [[("4. Объяснить", True)], ["разметка на снимке"], ["протокол и совет"]]]
for i, b in enumerate(boxes):
    x = x0 + (bw + gap) * i
    d.flow_box(s, x, yy, bw, bh, b, pt=17)
    if i:
        d.arrow(s, x - gap, yy + bh / 2, x, yy + bh / 2)
d.box(s, d.L, yy + bh + Inches(0.25), d.CW, Inches(1.2),
      "Постановщики задачи решали контроль укладки рентгенограмм грудной клетки ансамблем сетей на 69 000 снимков. "
      "При 1000 снимков и меньше опубликованные системы сначала измеряют анатомию, затем классифицируют.", 18)
d.refs(s, "Borisov AA, et al. Int J Comput Assist Radiol Surg. 2025;20(9):1829–1833. PMID: 40531386 · "
          "Meng Y, et al. Eur Radiol. 2022;32(11):7680–7690. PMID: 35420306 (1025 снимков, 10 индексов) · "
          "Ryu SM, et al. Med Phys. 2026;53(2):e70285. PMID: 41579106 (500 снимков, 14 ориентиров, ошибка 2,0 мм)")

s, y = d.slide("Один конвейер обслуживает командную строку, HTTP-сервис и контейнер", "Архитектура")
bw, bh, gap = Inches(2.1), Inches(1.55), Inches(0.26); yy = y + Inches(0.1)
boxes = [[[("DICOM, zip", True)], ["дубликаты — по хэшу"]], [[("Область", True)], ["по ширине кадра"], ["сторона бедра"]],
         [[("Модели", True)], ["ориентиры, признаки"], ["15 сетей, карта"]], [[("Критерии", True)], ["5 вероятностей"], ["пороги по F1"]],
         [[("Вердикт", True)], ["класс, типы"], ["«проверить»"]]]
for i, b in enumerate(boxes):
    x = d.L + (bw + gap) * i
    d.flow_box(s, x, yy, bw, bh, b, pt=15)
    if i:
        d.arrow(s, x - gap, yy + bh / 2, x, yy + bh / 2)
d.table(s, [["Выход", "Назначение"],
            ["results.xlsx / csv", "таблица ТЗ: строка на файл; лист «по исследованиям» — что переснять"],
            ["overlays/*.png", "дополнительная серия: ориентиры, ось, рамки предметов, вероятности"],
            ["sr/*_SR.dcm", "DICOM Basic Text SR: вердикт, протокол измерений, совет лаборанту"],
            ["HTTP API и страница", "POST /api/analyze, Swagger; контейнер без GPU и без сети"]],
        yy + bh + Inches(0.3), [3.0, 8.75], size=14, row_h_in=0.42)
d.refs(s, "Исходный код и документация: github.com/YODA02041994/lct2026-dxa-quality · прототип: docsemenov.ru/dxa")

figure_slide("Эксперименты", "Признаки по снимку и перенос обучения подняли ROC-AUC с 0,74 до 0,85",
             os.path.join(FIG, "progress.png"),
             "Каждый шаг проверен на снимках вне обучения. Замороженная сеть общего назначения (DINOv2) давала 0,60–0,67.",
             "Собственные данные: docs/05_Эксперименты.md, docs/08, эксперименты 01–16; 249 снимков, кросс-валидация по исследованиям", 3.2)

figure_slide("Метрики", "По типам нарушений ROC-AUC от 0,80 до 0,92; при 6–10 примерах интервалы широкие",
             os.path.join(FIG, "forest.png"),
             "F1 по типам: 0,67 · 0,40 · 0,78 · 0,75 · 0,67; macro-F1 0,653.",
             "Собственные данные: docs/metrics.md; 95 % ДИ — бутстреп по исследованиям, 1000 повторов", 3.2)

figure_slide("Метрики", "Бинарная классификация: ROC-AUC 0,849, F1 0,762, accuracy 0,859",
             os.path.join(FIG, "binary.png"),
             "Модель помечает 30 % снимков при 29 % нарушений в разметке экспертов.",
             "Собственные данные: docs/metrics.md; 249 снимков вне обучения; 95 % ДИ — бутстреп по исследованиям", 3.3)

figure_slide("Перенос обучения", "Предобучение на 9000 рентгенограмм грудной клетки повысило ROC-AUC сети по предметам с 0,78 до 0,86",
             os.path.join(FIG, "transfer.png"),
             "Карта предметов, обученная на тех же рентгенограммах, находит предметы на денситометрии без дообучения: ROC-AUC 0,84.",
             "Набор object-CXR (JF Healthcare, MIDL 2020; CC BY-NC 4.0): 10 000 рентгенограмм, 5000 с посторонними предметами · "
             "RadImageNet. Radiology: Artificial Intelligence. DOI: 10.1148/ryai.210315", 2.8)

figure_slide("Сравнение сетей", "Увеличение сети не повысило ROC-AUC; предобучение на рентгенограммах повысило с 0,82 до 0,86",
             os.path.join(FIG, "arch.png"),
             "Три сети DenseNet121 с рентгенограмм в ансамбле: ROC-AUC критерия 0,866 → 0,873, разница +0,007 [−0,003; +0,019]. Сети включены в модель.",
             "Собственные данные: experiments/exp16_architectures.py, docs/08_Рисерч_что_ещё.md; 25 эпох, одинаковые разбиения по исследованиям · "
             "веса DenseNet121: Cohen JP, et al. TorchXRayVision: a library of chest X-ray datasets and models. arXiv:2111.00595 (2021)", 2.9)

s, y = d.slide("На вырезке бедра сеть опирается на область малого вертела — признак ротации по стандарту ISCD", "Объяснимость")
ph = d.pic(s, os.path.join(IMG, "gradcam.jpg"), y, max_w=d.CW, max_h=Inches(3.1))
d.box(s, d.L, d.below(ph, 0.1), d.CW, Inches(0.7), "Верхний ряд — снимки с нарушением укладки (эксперт), нижний — норма. Карты Grad-CAM, снимки обучающего набора.", 18)
d.refs(s, "ISCD Official Positions 2023: бедро ротировано внутрь, малый вертел виден минимально или не виден · "
          "Собственные данные: scripts/gradcam_examples.py")

figure_slide("Анализ ошибок", "Из 16 проверенных подходов в модель вошли семь",
             os.path.join(FIG, "tried.png"),
             "Оранжевое — вошло в модель, серое — проверено и отклонено. 45 снимков с расхождением модели и эксперта вынесены на сверку.",
             "Собственные данные: docs/08_Рисерч_что_ещё.md, docs/09_Спорные_метки.md · RAD-DINO: arXiv:2311.13668", 3.15)

figure_slide("Порог", "При пороге сдачи чувствительность 0,78 и специфичность 0,89; помечается 30 % снимков",
             os.path.join(FIG, "tradeoff.png"),
             "Клинический режим (--mode sensitive): чувствительность 0,94, специфичность 0,38. Снимки у порога получают отметку «проверить» (16 %).",
             "Собственные данные: docs/08_Рисерч_что_ещё.md, раздел 6 · цена ошибок: Krueger D, et al. J Clin Densitom. 2019;22(1):115–124. PMID: 30327243", 3.0)

s, y = d.slide("Исследование из трёх снимков обрабатывается за 7–20 секунд на двух ядрах без GPU", "Скорость")
d.table(s, [["Конфигурация", "Измерение", "Время"],
            ["Сервер, 2 ядра CPU, 9 ГБ ОЗУ, без GPU", "3 снимка, полный цикл через HTTP", "7–20 с"],
            ["Контейнер linux/amd64, CPU, без сети", "3 снимка, командная строка", "6–15 с"],
            ["Apple M3 Max, GPU (MPS)", "499 файлов, 252 уникальных снимка", "28 с"],
            ["Требование ТЗ", "одно исследование", "≤ 180 с"]],
        y + Inches(0.05), [5.2, 4.6, 1.95], size=16, row_h_in=0.52)
d.box(s, d.L, y + Inches(2.95), d.CW, Inches(0.9),
      "Минимум: 2 ядра, 6 ГБ ОЗУ. Рекомендуется: 8 ядер, 16 ГБ ОЗУ. Архив образа 1,4 ГБ, веса 0,9 ГБ, версии зависимостей зафиксированы.", 18)
d.refs(s, "Собственные замеры 24–28.09.2026: README.md, раздел «Железо и время»; requirements.lock")

s, y = d.slide("Кейсы: посторонний предмет, неверная укладка бедра, норма", "Кейсы")
d.pic(s, os.path.join(IMG, "cases.jpg"), y, max_w=d.CW, max_h=Inches(3.95))
d.refs(s, "Собственные данные: снимки вне обучения; рядом с вердиктом — оценка эксперта постановщика; scripts/export_examples.py")

s, y = d.slide("Лаборант получает вердикт, протокол измерений и совет по исправлению укладки", "Прототип")
d.pic(s, os.path.join(IMG, "ui.png"), y, max_w=d.CW, max_h=Inches(3.95))
d.refs(s, "Прототип: https://docsemenov.ru/dxa/ · описание API: https://docsemenov.ru/dxa/docs")

s, y = d.slide("Ограничения", "Ограничения")
d.bullets(s, ["Обучение и проверка — один аппарат (GE Lunar Prodigy Advance); на других аппаратах нужна проверка.",
              "Для трёх критериев в данных 6–10 снимков с нарушением: интервалы ROC-AUC шириной до 0,5.",
              "Ось позвоночника: у снимков с нарушением наклон 2,6–7,2°, у нормы до 9°; порог 5° разметкой не подтверждается.",
              "45 снимков с расхождением модели и эксперта требуют сверки; до неё метрики считаются по исходной разметке."], y + Inches(0.05), size=19)
d.refs(s, "Собственные данные: docs/02_Данные_аудит.md, docs/09_Спорные_метки.md")

s, y = d.slide("Сценарий внедрения: от аудита архива к подсказке лаборанту", "Внедрение")
bw, bh, gap = Inches(3.7), Inches(1.9), Inches(0.37); yy = y + Inches(0.15)
boxes = [[[("1. Аудит архива", True)], ["сортировка исследований"], ["по вероятности нарушения"]],
         [[("2. Контроль на потоке", True)], ["допуск к расчёту плотности"], ["только годных снимков"]],
         [[("3. Подсказка лаборанту", True)], ["совет до ухода пациента,"], ["пересъёмка при нарушении"]]]
for i, b in enumerate(boxes):
    x = d.L + (bw + gap) * i
    d.flow_box(s, x, yy, bw, bh, b, pt=18)
    if i:
        d.arrow(s, x - gap, yy + bh / 2, x, yy + bh / 2)
d.box(s, d.L, yy + bh + Inches(0.3), d.CW, Inches(1.0),
      "Каждый шаг использует один и тот же контейнер. Размеченные врачом отказы возвращаются в обучение.", 18)
d.refs(s, "Сценарии применения систем контроля качества: Borisov AA, et al. Int J Comput Assist Radiol Surg. 2025;20(9):1829–1833. PMID: 40531386")

s, y = d.slide("Материалы решения", "Ссылки")
d.table(s, [["Материал", "Адрес"],
            ["Исходный код, README", "github.com/YODA02041994/lct2026-dxa-quality"],
            ["Прототип и описание API", "docsemenov.ru/dxa · docsemenov.ru/dxa/docs"],
            ["Веса моделей (18 файлов)", "github.com/YODA02041994/lct2026-dxa-quality/releases"],
            ["Документация", "docs/: руководство пользователя, развёртывание, обучение"],
            ["Демонстрация", "docs/DEMO.md · запись: docsemenov.ru/dxa/files/dxaqc_demo.gif"],
            ["Журнал экспериментов", "docs/05_Эксперименты.md, docs/08_Рисерч_что_ещё.md"]],
        y + Inches(0.05), [4.2, 7.55], size=16, row_h_in=0.5)
d.refs(s, "Команда «" + T("team_name", "[название команды]") + "», ЛЦТ-2026, задача № 4")

d.save(OUT)
os.remove(tmp)
print("сохранено:", OUT, "| слайдов:", len(d.prs.slides))
