"""Чтение экспертной разметки `разметка.xlsx` (лист «Калибровка») и официальный словарь меток.

Структура листа: 2 строки шапки, далее строка = исследование.
Колонки (0-based): 0 №, 1 study (имя папки), 2..4 позвоночник, 5..6 правое бедро,
7..8 левое бедро, 9..11 Итог по областям, 12 комментарий, 14..18 сводный блок (игнорируем).

Словарь значений — из документа организатора «Разъяснения по вопросам ЛЦТ_V2.docx» (16.09.2026)
с поправкой модератора в чате (16.09 14:01): «Не выровнена ось позвоночника», не «выравнена».
"""
from __future__ import annotations

from dataclasses import dataclass

import openpyxl

# ---- Официальные значения выходной таблицы (организатор, 16.09.2026) ----------------------
REGION_SPINE = "Поясничный отдел позвоночника"
REGION_HIP = "Проксимальный отдел бедра"          # сторона (лево/право) организатору не важна
VIOLATION_SEPARATOR = ";"                          # несколько нарушений — через «;», нет нарушений — пусто
PROB_COLUMN = "quality_prob"                       # вероятность нарушения в [0;1], разрешена организатором

# внутренний код → официальный текст violation_type
SPINE_FLAGS = ["spine_positioning", "spine_axis_tilt", "spine_artifact"]
HIP_FLAGS = ["hip_positioning_rotation", "hip_roi_field"]

VIOLATION_TEXT = {
    "spine_positioning": "Некорректная укладка",
    "spine_axis_tilt": "Не выровнена ось позвоночника",
    "spine_artifact": "Присутствуют посторонние предметы",
    "hip_positioning_rotation": "Некорректная укладка",
    "hip_roi_field": "Некорректная область интереса",
}

SPINE_FLAGS_RU = {
    "spine_positioning": "некорректная укладка (Th12 / гребни подвздошных костей)",
    "spine_axis_tilt": "наклон оси позвоночника > 5°",
    "spine_artifact": "посторонние предметы / артефакты",
}
HIP_FLAGS_RU = {
    "hip_positioning_rotation": "позиционирование / ротация бедра",
    "hip_roi_field": "недостаточное поле области интереса",
}

# Размер пикселя. Организатор: сканер 0,6 мм (X) × 1,05 мм (Y) — это нативный шаг строк.
# Проверено 16.09 (docs/02, раздел «Масштаб»): выгрузка пересэмплирована в КВАДРАТНЫЙ пиксель —
# ExposedArea, шаг позвонков и круглая головка бедра сходятся на ≈0,6 мм/px по обеим осям.
PIXEL_MM_X_SCANNER = 0.6
PIXEL_MM_Y_SCANNER = 1.05
PIXEL_MM_EXPORT = {300: 0.60, 280: 0.65, 248: 0.61}   # по ширине изображения → мм/px (обе оси), ±10 %


def pixel_mm(cols: int) -> float:
    """мм на пиксель для выгруженного изображения по его ширине; неизвестная ширина → 0,6."""
    return PIXEL_MM_EXPORT.get(int(cols), 0.6)


def violation_string(flags: dict[str, int | float | bool]) -> str:
    """Собрать `violation_type` из положительных флагов по официальному словарю."""
    texts: list[str] = []
    for code, val in flags.items():
        if val and VIOLATION_TEXT.get(code) and VIOLATION_TEXT[code] not in texts:
            texts.append(VIOLATION_TEXT[code])
    return VIOLATION_SEPARATOR.join(texts)


@dataclass
class RegionLabel:
    flags: dict[str, int | None]
    total: int | None  # колонка «Итог» эксперта

    @property
    def labeled(self) -> bool:
        return self.total is not None or any(v is not None for v in self.flags.values())

    @property
    def flags_positive(self) -> bool | None:
        vals = [v for v in self.flags.values() if v is not None]
        if not vals:
            return None
        return any(v == 1 for v in vals)

    @property
    def inconsistent(self) -> bool:
        """Итог эксперта расходится с флагами (в данных 3 таких случая)."""
        fp = self.flags_positive
        return fp is not None and self.total is not None and int(fp) != self.total

    @property
    def truth(self) -> int | None:
        """Истинная метка изображения по правилу организатора (сессия 16.09):
        верить колонкам критериев, НЕ колонке «Итог»; любой флаг = 1 → нарушение.
        None — область не размечена (в исследовании её нет)."""
        fp = self.flags_positive
        return None if fp is None else int(fp)

    @property
    def violation_type(self) -> str:
        return violation_string({k: (v == 1) for k, v in self.flags.items()})


@dataclass
class StudyLabel:
    num: int
    study_dir: str
    spine: RegionLabel
    hip_r: RegionLabel
    hip_l: RegionLabel
    comment: str | None

    @property
    def any_violation(self) -> bool:
        """По правилу организатора: из флагов, не из «Итог»."""
        return any(r.truth == 1 for r in (self.spine, self.hip_r, self.hip_l))


def _int_or_none(v):
    if v is None or v == "":
        return None
    return int(v)


def read_labels(xlsx_path: str) -> list[StudyLabel]:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb["Калибровка"] if "Калибровка" in wb.sheetnames else wb.worksheets[0]
    out: list[StudyLabel] = []
    for row in list(ws.iter_rows(values_only=True))[2:]:
        if not row[1]:
            continue
        spine = RegionLabel(
            flags=dict(zip(SPINE_FLAGS, map(_int_or_none, row[2:5]))), total=_int_or_none(row[9])
        )
        hip_r = RegionLabel(
            flags=dict(zip(HIP_FLAGS, map(_int_or_none, row[5:7]))), total=_int_or_none(row[10])
        )
        hip_l = RegionLabel(
            flags=dict(zip(HIP_FLAGS, map(_int_or_none, row[7:9]))), total=_int_or_none(row[11])
        )
        out.append(
            StudyLabel(
                num=int(row[0]) if row[0] is not None else len(out) + 1,
                study_dir=str(row[1]).strip(),
                spine=spine,
                hip_r=hip_r,
                hip_l=hip_l,
                comment=(str(row[12]).strip() if row[12] else None),
            )
        )
    return out
